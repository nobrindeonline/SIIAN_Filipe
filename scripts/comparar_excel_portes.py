import os
import sys
from openpyxl import load_workbook


BASE_DIR = os.path.dirname(
    os.path.dirname(
        os.path.abspath(__file__)
    )
)

if BASE_DIR not in sys.path:
    sys.path.insert(0, BASE_DIR)


from config.env_settings import get_paths


def obter_headers(ws):
    return {
        str(ws.cell(row=1, column=coluna).value).strip(): coluna
        for coluna in range(1, ws.max_column + 1)
        if ws.cell(row=1, column=coluna).value is not None
    }


def comparar_folha(
    ws_principal,
    ws_importacao,
    coluna_ref
):
    headers_principal = obter_headers(ws_principal)
    headers_importacao = obter_headers(ws_importacao)

    mapa_principal = {}

    for linha in range(2, ws_principal.max_row + 1):
        ref = ws_principal.cell(
            row=linha,
            column=headers_principal[coluna_ref]
        ).value

        if ref is not None:
            mapa_principal[str(ref).strip()] = linha

    diferencas = 0

    for linha_import in range(
        2,
        ws_importacao.max_row + 1
    ):
        ref = ws_importacao.cell(
            row=linha_import,
            column=headers_importacao[coluna_ref]
        ).value

        if ref is None:
            continue

        ref = str(ref).strip()

        if ref not in mapa_principal:
            print(f"NOVA REFERENCIA: {ref}")
            continue

        linha_principal = mapa_principal[ref]

        for coluna, coluna_import in headers_importacao.items():

            if coluna not in headers_principal:
                continue

            if coluna in (
                "id_portes",
                "id_zona_geo",
                "data_update"
            ):
                continue

            valor_import = ws_importacao.cell(
                row=linha_import,
                column=coluna_import
            ).value

            valor_principal = ws_principal.cell(
                row=linha_principal,
                column=headers_principal[coluna]
            ).value

            if valor_import != valor_principal:
                diferencas += 1

                print("")
                print(f"REFERENCIA: {ref}")
                print(f"COLUNA: {coluna}")
                print(
                    f"PRINCIPAL: {repr(valor_principal)}"
                )
                print(
                    f"IMPORTACAO: {repr(valor_import)}"
                )

    return diferencas


def run():
    paths = get_paths()

    wb_principal = load_workbook(
        paths["EXCEL_PORTES_ENVIO"],
        data_only=False
    )

    wb_importacao = load_workbook(
        paths["EXCEL_PORTES_IMPORTACAO"],
        data_only=False
    )

    print("=== PORTES ===")

    diferencas_portes = comparar_folha(
        wb_principal["portes_envio"],
        wb_importacao["portes_envio"],
        "ref_portes"
    )

    print("")
    print("=== ZONAS ===")

    diferencas_zonas = comparar_folha(
        wb_principal["zonas_geograficas"],
        wb_importacao["zonas_geograficas"],
        "ref_zona_geo"
    )

    print("")
    print("==============================")
    print(
        f"Diferencas portes: {diferencas_portes}"
    )
    print(
        f"Diferencas zonas: {diferencas_zonas}"
    )
    print("==============================")


if __name__ == "__main__":
    run()
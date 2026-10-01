import os
import sys
import shutil
from datetime import date, datetime

from openpyxl import load_workbook


BASE_DIR = os.path.dirname(
    os.path.dirname(
        os.path.abspath(__file__)
    )
)

if BASE_DIR not in sys.path:
    sys.path.insert(0, BASE_DIR)


from config.env_settings import get_paths


def obter_headers(ws):
    headers = {}

    for coluna in range(1, ws.max_column + 1):
        valor = ws.cell(
            row=1,
            column=coluna
        ).value

        if valor is not None:
            headers[str(valor).strip()] = coluna

    return headers


def valores_iguais(valor1, valor2):
    if valor1 is None and valor2 is None:
        return True

    if valor1 == "" and valor2 is None:
        return True

    if valor2 == "" and valor1 is None:
        return True

    return valor1 == valor2


def validar_referencias(
    ws,
    headers,
    coluna_ref
):
    coluna = headers[coluna_ref]

    referencias = set()
    duplicadas = set()

    for linha in range(2, ws.max_row + 1):
        ref = ws.cell(
            row=linha,
            column=coluna
        ).value

        if ref is None:
            continue

        ref = str(ref).strip()

        if not ref:
            continue

        if ref in referencias:
            duplicadas.add(ref)

        referencias.add(ref)

    if duplicadas:
        raise ValueError(
            f"Referencias duplicadas: "
            f"{sorted(duplicadas)}"
        )


def importar_folha(
    ws_principal,
    ws_importacao,
    coluna_ref,
    coluna_id
):
    headers_principal = obter_headers(
        ws_principal
    )

    headers_importacao = obter_headers(
        ws_importacao
    )

    if coluna_ref not in headers_principal:
        raise ValueError(
            f"Coluna '{coluna_ref}' nao existe "
            f"no ficheiro principal."
        )

    if coluna_ref not in headers_importacao:
        raise ValueError(
            f"Coluna '{coluna_ref}' nao existe "
            f"no ficheiro de importacao."
        )

    validar_referencias(
        ws_importacao,
        headers_importacao,
        coluna_ref
    )

    coluna_ref_principal = (
        headers_principal[coluna_ref]
    )

    coluna_ref_importacao = (
        headers_importacao[coluna_ref]
    )

    mapa_principal = {}

    for linha in range(
        2,
        ws_principal.max_row + 1
    ):
        ref = ws_principal.cell(
            row=linha,
            column=coluna_ref_principal
        ).value

        if ref is None:
            continue

        ref = str(ref).strip()

        if ref:
            mapa_principal[ref] = linha

    proximo_id = 1

    if coluna_id in headers_principal:
        ids = []

        coluna_id_principal = (
            headers_principal[coluna_id]
        )

        for linha in range(
            2,
            ws_principal.max_row + 1
        ):
            valor = ws_principal.cell(
                row=linha,
                column=coluna_id_principal
            ).value

            if isinstance(valor, (int, float)):
                ids.append(int(valor))

        if ids:
            proximo_id = max(ids) + 1

    inseridos = 0
    atualizados = 0
    sem_alteracao = 0

    for linha_import in range(
        2,
        ws_importacao.max_row + 1
    ):
        ref = ws_importacao.cell(
            row=linha_import,
            column=coluna_ref_importacao
        ).value

        if ref is None:
            continue

        ref = str(ref).strip()

        if not ref:
            continue

        if ref in mapa_principal:
            linha_principal = mapa_principal[ref]

            houve_alteracao = False

            for nome_coluna, coluna_import in (
                headers_importacao.items()
            ):

                if nome_coluna not in headers_principal:
                    continue

                if nome_coluna in (
                    coluna_id,
                    coluna_ref,
                    "data_update"
                ):
                    continue

                coluna_principal = (
                    headers_principal[nome_coluna]
                )

                valor_import = ws_importacao.cell(
                    row=linha_import,
                    column=coluna_import
                ).value

                valor_atual = ws_principal.cell(
                    row=linha_principal,
                    column=coluna_principal
                ).value

                if not valores_iguais(
                    valor_import,
                    valor_atual
                ):
                    ws_principal.cell(
                        row=linha_principal,
                        column=coluna_principal
                    ).value = valor_import

                    houve_alteracao = True

            if houve_alteracao:
                if "data_update" in headers_principal:
                    ws_principal.cell(
                        row=linha_principal,
                        column=headers_principal[
                            "data_update"
                        ]
                    ).value = date.today()

                atualizados += 1

            else:
                sem_alteracao += 1

        else:
            nova_linha = (
                ws_principal.max_row + 1
            )

            for nome_coluna, coluna_principal in (
                headers_principal.items()
            ):

                if nome_coluna == coluna_id:
                    ws_principal.cell(
                        row=nova_linha,
                        column=coluna_principal
                    ).value = proximo_id
                    continue

                if nome_coluna == "data_update":
                    ws_principal.cell(
                        row=nova_linha,
                        column=coluna_principal
                    ).value = date.today()
                    continue

                if nome_coluna in headers_importacao:
                    valor = ws_importacao.cell(
                        row=linha_import,
                        column=headers_importacao[
                            nome_coluna
                        ]
                    ).value

                    ws_principal.cell(
                        row=nova_linha,
                        column=coluna_principal
                    ).value = valor

            mapa_principal[ref] = nova_linha

            proximo_id += 1
            inseridos += 1

    return (
        inseridos,
        atualizados,
        sem_alteracao
    )


def criar_backup(ficheiro_principal):
    pasta_data = os.path.dirname(
        ficheiro_principal
    )

    pasta_backup = os.path.join(
        pasta_data,
        "backup"
    )

    os.makedirs(
        pasta_backup,
        exist_ok=True
    )

    timestamp = datetime.now().strftime(
        "%Y%m%d_%H%M%S"
    )

    caminho_backup = os.path.join(
        pasta_backup,
        f"portes_envio_{timestamp}.xlsx"
    )

    shutil.copy2(
        ficheiro_principal,
        caminho_backup
    )

    return caminho_backup


def run():
    print("")
    print(
        "========================================"
    )
    print(" IMPORTACAO EXCEL DE PORTES")
    print(
        "========================================"
    )
    print("")

    paths = get_paths()

    ficheiro_principal = (
        paths["EXCEL_PORTES_ENVIO"]
    )

    ficheiro_importacao = (
        paths["EXCEL_PORTES_IMPORTACAO"]
    )

    if not os.path.exists(
        ficheiro_principal
    ):
        raise FileNotFoundError(
            f"Ficheiro principal nao encontrado: "
            f"{ficheiro_principal}"
        )

    if not os.path.exists(
        ficheiro_importacao
    ):
        raise FileNotFoundError(
            f"Ficheiro de importacao nao encontrado: "
            f"{ficheiro_importacao}"
        )

    print(
        f"Principal: "
        f"{ficheiro_principal}"
    )

    print(
        f"Importacao: "
        f"{ficheiro_importacao}"
    )

    wb_principal = load_workbook(
        ficheiro_principal
    )

    wb_importacao = load_workbook(
        ficheiro_importacao
    )

    for folha in [
        "portes_envio",
        "zonas_geograficas"
    ]:
        if folha not in wb_principal.sheetnames:
            raise ValueError(
                f"Folha '{folha}' nao existe "
                f"no ficheiro principal."
            )

        if folha not in wb_importacao.sheetnames:
            raise ValueError(
                f"Folha '{folha}' nao existe "
                f"no ficheiro de importacao."
            )

    print("A verificar portes...")

    (
        portes_inseridos,
        portes_atualizados,
        portes_iguais
    ) = importar_folha(
        wb_principal["portes_envio"],
        wb_importacao["portes_envio"],
        "ref_portes",
        "id_portes"
    )

    print(
        "A verificar zonas geograficas..."
    )

    (
        zonas_inseridas,
        zonas_atualizadas,
        zonas_iguais
    ) = importar_folha(
        wb_principal["zonas_geograficas"],
        wb_importacao["zonas_geograficas"],
        "ref_zona_geo",
        "id_zona_geo"
    )

    caminho_backup = criar_backup(
        ficheiro_principal
    )

    wb_principal.save(
        ficheiro_principal
    )

    print("")
    print(
        "========================================"
    )
    print(" IMPORTACAO CONCLUIDA")
    print(
        "========================================"
    )
    print("")

    print("PORTES")
    print(
        f"  Inseridos: "
        f"{portes_inseridos}"
    )
    print(
        f"  Atualizados: "
        f"{portes_atualizados}"
    )
    print(
        f"  Sem alteracao: "
        f"{portes_iguais}"
    )

    print("")
    print("ZONAS")
    print(
        f"  Inseridas: "
        f"{zonas_inseridas}"
    )
    print(
        f"  Atualizadas: "
        f"{zonas_atualizadas}"
    )
    print(
        f"  Sem alteracao: "
        f"{zonas_iguais}"
    )

    print("")
    print(
        f"Backup criado: "
        f"{caminho_backup}"
    )

    print("")
    print(
        "Ficheiro principal atualizado."
    )


if __name__ == "__main__":
    run()